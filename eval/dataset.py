from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from eval.golden import GoldenItem


_QUESTION_HEADERS = {"질문", "question", "query"}
_GROUND_TRUTH_HEADERS = {
    "정답",
    "모범답변",
    "기대답변",
    "답변",
    "ground_truth",
    "reference_answer",
}
_ID_HEADERS = {"순번", "id", "번호", "문항번호"}
_QTYPE_HEADERS = {"유형", "질문유형", "qtype", "category"}


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _is_blank_row(row: Iterable[Any]) -> bool:
    return all(_is_blank(value) for value in row)


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header(value: Any) -> str:
    return _as_text(value).casefold()


def _column_index(headers: list[str], aliases: set[str]) -> int | None:
    return next((index for index, header in enumerate(headers) if header in aliases), None)


def _cell(row: tuple[Any, ...], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def resolve_workbook(title: str, root: Path = Path("Docs")) -> Path:
    safe_title = title.strip()
    if (
        not safe_title
        or safe_title in {".", ".."}
        or "/" in safe_title
        or "\\" in safe_title
        or Path(safe_title).suffix
    ):
        raise ValueError("데이터셋 제목은 확장자와 경로가 없는 파일명이어야 합니다")

    workbook = root / f"{safe_title}.xlsx"
    try:
        workbook.resolve().relative_to(root.resolve())
    except ValueError as error:
        raise ValueError("평가 데이터셋은 Docs 루트 안에 있어야 합니다") from error
    if not workbook.is_file():
        raise FileNotFoundError(f"평가 데이터셋을 찾을 수 없습니다: {workbook}")
    return workbook


def load_golden_workbook(path: Path) -> list[GoldenItem]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        if worksheet is None:
            raise ValueError("활성 워크시트를 찾을 수 없습니다")

        rows = enumerate(worksheet.iter_rows(values_only=True), start=1)
        header_row_number: int | None = None
        header_values: tuple[Any, ...] | None = None
        for row_number, row in rows:
            if not _is_blank_row(row):
                header_row_number = row_number
                header_values = row
                break

        if header_values is None or header_row_number is None:
            raise ValueError("헤더 행을 찾을 수 없습니다")

        headers = [_header(value) for value in header_values]
        question_index = _column_index(headers, _QUESTION_HEADERS)
        ground_truth_index = _column_index(headers, _GROUND_TRUTH_HEADERS)
        id_index = _column_index(headers, _ID_HEADERS)
        qtype_index = _column_index(headers, _QTYPE_HEADERS)

        missing_columns: list[str] = []
        if question_index is None:
            missing_columns.append("질문")
        if ground_truth_index is None:
            missing_columns.append("정답")
        if missing_columns:
            names = ", ".join(missing_columns)
            raise ValueError(
                f"시트 '{worksheet.title}'의 헤더 행 {header_row_number}에 "
                f"필수 열이 없습니다: {names}"
            )

        items: list[GoldenItem] = []
        id_rows: dict[str, int] = {}
        question_rows: dict[str, int] = {}
        for row_number, row in rows:
            if _is_blank_row(row):
                continue

            question = _as_text(_cell(row, question_index))
            if not question:
                raise ValueError(f"Excel 행 {row_number}의 질문 값이 비어 있습니다")

            ground_truth = _as_text(_cell(row, ground_truth_index))
            if not ground_truth:
                raise ValueError(f"Excel 행 {row_number}의 정답 값이 비어 있습니다")

            item_id = _as_text(_cell(row, id_index)) or f"row-{row_number}"
            qtype = _as_text(_cell(row, qtype_index)) or "미분류"

            if item_id in id_rows:
                first_row = id_rows[item_id]
                raise ValueError(
                    f"중복 ID '{item_id}': Excel 행 {first_row}, {row_number}"
                )
            if question in question_rows:
                first_row = question_rows[question]
                raise ValueError(
                    f"중복 질문 '{question}': Excel 행 {first_row}, {row_number}"
                )

            id_rows[item_id] = row_number
            question_rows[question] = row_number
            items.append(
                GoldenItem(
                    id=item_id,
                    question=question,
                    qtype=qtype,
                    ground_truth=ground_truth,
                    sheet=worksheet.title,
                    row_number=row_number,
                )
            )

        if not items:
            raise ValueError(f"시트 '{worksheet.title}'에 평가 문항이 없습니다")
        return items
    finally:
        workbook.close()